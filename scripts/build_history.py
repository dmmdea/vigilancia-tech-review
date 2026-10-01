#!/usr/bin/env python3
"""Per-student history of earlier rounds -> the dates a tool's age may be
measured from (the "fecha ancla").

Teaching-team rule (2026-10-01, final round): carrying the SAME tool from one
round to the next is valid — each new submission is an update of the same
case. So a tool's age is measured at the FIRST round in which the student
presented it, not at the date of the latest submission, and NEVER at the date
the grading happens to run (grading days after the deadline would otherwise
disqualify students for the calendar, not for their choice). A tool that is
new in this submission is measured at the student's own submission date.

This script only gathers facts; the reviewer decides whether the capability
shown NOW is the one presented then, and validate_review.py checks the
arithmetic against the dates listed here.

Usage:
    python build_history.py <workdir> --master=<xlsx> [--master=<xlsx>]...
        [--round-date="<ronda>=YYYY-MM-DD"]...   # when a round has no
                                                 # 'Meta - <ronda>' sheet
        [--fecha-entrega-defecto=YYYY-MM-DD]     # folders whose Canvas
                                                 # timestamp cannot be read

Reads <workdir>/review_plan.json and every 'Ranking - <ronda>' sheet of the
masters (delivered results; graded rows only, Estado = REVISADO). A student
is matched by the stable 'Clave' (Canvas student id). Rows of a round that
carries no Clave are matched by EXACT normalized name and labelled as such —
never fuzzy: two students can share a name, and a wrong match moves a grade.
A round's date comes from --round-date, else from 'Meta - <ronda>' (Fecha de
corrida); a round with neither stops the script (exit 2) — a guessed date
would silently move every anchor in it.

Writes <workdir>/historial.json:
  {"rondas": [{"ronda", "fecha", "fuente"}],
   "folders": {<folder_id>: {"student_name", "canvas_key", "fecha_entrega",
               "rondas": [{"ronda", "fecha", "herramienta",
                           "fecha_verificada", "descalificado",
                           "coincidencia": "clave" | "nombre"}],
               "anclas_permitidas": ["YYYY-MM-DD", ...],
               "historial_block": "<texto para el revisor>"}}}
Exit: 0 ok · 1 usage · 2 unreadable input / undated round / unreadable
submission date.
"""
import json
import os
import re
import sys
import unicodedata

for _s in (sys.stdout, sys.stderr):
    if _s in (sys.__stdout__, sys.__stderr__) and hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8")

MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5,
         "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9,
         "octubre": 10, "noviembre": 11, "diciembre": 12}
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def norm_name(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return " ".join(s.casefold().split())


def norm_key(v):
    """Excel may hand back 11717, 11717.0 or '11717 ' for the same Clave."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def iso_date(v):
    if v is None:
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()[:10]
    return s if ISO.match(s) else None


def submission_date(timestamp):
    """'24 de septiembre de 2026 1841' -> '2026-09-24' (None if unreadable)."""
    m = re.search(r"(\d{1,2})\s+de\s+([A-Za-záéíóúñ]+)\s+de\s+(\d{4})",
                  timestamp or "", re.IGNORECASE)
    if not m or m.group(2).lower() not in MESES:
        return None
    return f"{int(m.group(3)):04d}-{MESES[m.group(2).lower()]:02d}-{int(m.group(1)):02d}"


def load_rounds(paths, overrides):
    """[(ronda, fecha, fuente, rows)] from every 'Ranking - X' sheet."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: falta openpyxl (pip install -r requirements.txt).",
              file=sys.stderr)
        sys.exit(2)
    rounds, seen, undated = [], {}, []
    for p in paths:
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
        except Exception as e:
            print(f"ERROR: no se pudo abrir {p}: {e}", file=sys.stderr)
            sys.exit(2)
        for ws in wb.worksheets:
            if not ws.title.startswith("Ranking - "):
                continue
            ronda = ws.title[len("Ranking - "):].strip()
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
            need = ("Estudiante", "Herramienta", "Estado")
            if not all(h in hdr for h in need):
                print(f"AVISO: {os.path.basename(p)} › {ws.title} no trae "
                      f"{need} — hoja ignorada.", file=sys.stderr)
                continue
            col = {h: hdr.index(h) for h in hdr if h}

            def cell(r, h):
                i = col.get(h)
                return r[i] if i is not None and i < len(r) else None
            graded = []
            for r in rows[1:]:
                if str(cell(r, "Estado") or "").strip().upper() != "REVISADO":
                    continue
                graded.append({
                    "clave": norm_key(cell(r, "Clave")),
                    "nombre": norm_name(cell(r, "Estudiante")),
                    "herramienta": str(cell(r, "Herramienta") or "").strip(),
                    "fecha_verificada": iso_date(cell(r, "Fecha lanz. verificada"))
                    or str(cell(r, "Fecha lanz. verificada") or "").strip()[:10],
                    "descalificado": str(cell(r, "¿Descalificado?") or "")
                    .strip().upper().startswith("S"),
                })
            if ronda in seen:
                if len(graded) != seen[ronda]:
                    print(f"AVISO: la ronda '{ronda}' aparece en más de un "
                          f"maestro con distinto número de filas "
                          f"({seen[ronda]} vs {len(graded)}); se usa la "
                          "primera.", file=sys.stderr)
                continue
            fecha = overrides.get(ronda)
            if not fecha and f"Meta - {ronda}" in wb.sheetnames:
                meta = wb[f"Meta - {ronda}"]
                for r in meta.iter_rows(values_only=True):
                    if r and str(r[0] or "").strip() == "Fecha de corrida":
                        fecha = iso_date(r[1] if len(r) > 1 else None)
                        break
            if not fecha:
                undated.append(f"{os.path.basename(p)} › {ws.title}")
                continue
            seen[ronda] = len(graded)
            rounds.append((ronda, fecha, f"{os.path.basename(p)} › {ws.title}",
                           graded))
    if undated:
        print("ERROR: rondas sin fecha (sin 'Meta - <ronda>' ni "
              "--round-date): " + "; ".join(undated) + ". Pasa "
              "--round-date=\"<ronda>=YYYY-MM-DD\" — una fecha adivinada "
              "movería todas las anclas de esa ronda.", file=sys.stderr)
        sys.exit(2)
    rounds.sort(key=lambda t: t[1])
    return rounds


def block_for(entry):
    lines = []
    if entry["rondas"]:
        lines.append("Rondas anteriores de este estudiante (resultados ya "
                     "entregados al equipo docente):")
        for r in entry["rondas"]:
            tag = ("" if r["coincidencia"] == "clave" else
                   " [coincidencia SOLO por nombre, sin código: confírmala "
                   "antes de usarla como ancla y dilo en evidence_notes]")
            lines.append(f"- {r['ronda']} ({r['fecha']}){tag}: herramienta "
                         f"«{r['herramienta']}»")
    else:
        lines.append("Sin rondas anteriores registradas para este estudiante "
                     "en los resultados entregados: toda herramienta es "
                     "nueva en esta entrega.")
    names = {r["fecha"]: r["ronda"] for r in entry["rondas"]}
    shown = []
    for d in entry["anclas_permitidas"]:
        what = ("fecha de entrega actual" if d == entry["fecha_entrega"]
                else names.get(d, ""))
        shown.append(f"{d} ({what})" if what else d)
    lines.append("Fechas ancla permitidas: " + " · ".join(shown) + ".")
    return "\n".join(lines)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    masters, overrides, default_date = [], {}, None
    for a in sys.argv[1:]:
        if a.startswith("--master="):
            masters.append(a.split("=", 1)[1])
        elif a.startswith("--round-date="):
            k, _, v = a.split("=", 1)[1].partition("=")
            if not ISO.match(v.strip()):
                print(f"ERROR: --round-date inválido: {a}", file=sys.stderr)
                sys.exit(1)
            overrides[k.strip()] = v.strip()
        elif a.startswith("--fecha-entrega-defecto="):
            default_date = a.split("=", 1)[1].strip()
            if not ISO.match(default_date):
                print(f"ERROR: fecha inválida: {a}", file=sys.stderr)
                sys.exit(1)
        elif a.startswith("--"):
            print(f"ERROR: opción desconocida: {a}", file=sys.stderr)
            sys.exit(1)
    if len(args) != 1 or not masters:
        print(__doc__, file=sys.stderr)
        sys.exit(1)
    work = args[0]
    try:
        with open(os.path.join(work, "review_plan.json"), encoding="utf-8") as f:
            plan = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        print(f"ERROR: review_plan.json ilegible: {e}", file=sys.stderr)
        sys.exit(2)

    rounds = load_rounds(masters, overrides)
    by_key, by_name = {}, {}
    for ronda, fecha, _src, rows in rounds:
        for r in rows:
            rec = {"ronda": ronda, "fecha": fecha,
                   "herramienta": r["herramienta"],
                   "fecha_verificada": r["fecha_verificada"],
                   "descalificado": r["descalificado"]}
            if r["clave"]:
                by_key.setdefault(r["clave"], []).append(dict(rec, coincidencia="clave"))
            elif r["nombre"]:
                by_name.setdefault(r["nombre"], []).append(dict(rec, coincidencia="nombre"))

    out, undated, stats = {}, [], {"clave": 0, "nombre": 0, "ninguna": 0}
    for e in plan["folders"]:
        fecha_entrega = submission_date(e.get("timestamp")) or default_date
        if not fecha_entrega:
            undated.append(e.get("folder_name", e["folder_id"]))
            continue
        key = norm_key(e.get("canvas_key"))
        hist = list(by_key.get(key, [])) if key else []
        keyed_rounds = {h["ronda"] for h in hist}
        # a name match only fills rounds the Clave did not already cover
        for h in by_name.get(norm_name(e.get("student_name")), []):
            if h["ronda"] not in keyed_rounds:
                hist.append(h)
        # a round dated AFTER this submission cannot be where it started
        hist = sorted((h for h in hist if h["fecha"] <= fecha_entrega),
                      key=lambda h: h["fecha"])
        anclas = sorted({h["fecha"] for h in hist} | {fecha_entrega})
        entry = {"student_name": e.get("student_name"),
                 "canvas_key": e.get("canvas_key"),
                 "fecha_entrega": fecha_entrega, "rondas": hist,
                 "anclas_permitidas": anclas}
        entry["historial_block"] = block_for(entry)
        out[e["folder_id"]] = entry
        kinds = {h["coincidencia"] for h in hist}
        stats["clave" if "clave" in kinds else
              "nombre" if kinds else "ninguna"] += 1
    if undated:
        print("ERROR: carpetas sin fecha de entrega legible (pasa "
              "--fecha-entrega-defecto=YYYY-MM-DD si corresponde): "
              + "; ".join(undated), file=sys.stderr)
        sys.exit(2)

    with open(os.path.join(work, "historial.json"), "w", encoding="utf-8") as f:
        json.dump({"rondas": [{"ronda": r, "fecha": d, "fuente": s}
                              for r, d, s, _rows in rounds],
                   "folders": out}, f, ensure_ascii=False, indent=2)
    print("rondas: " + ", ".join(f"{r} = {d}" for r, d, _s, _x in rounds))
    print(f"historial.json: {len(out)} carpetas | con historial por Clave: "
          f"{stats['clave']} | solo por nombre: {stats['nombre']} | sin "
          f"historial: {stats['ninguna']}")
    for fid, entry in out.items():
        if not any(h["coincidencia"] == "clave" for h in entry["rondas"]):
            how = ("solo por NOMBRE" if entry["rondas"] else "SIN historial")
            print(f"  {how}: {entry['student_name']} ({fid})")


if __name__ == "__main__":
    main()
